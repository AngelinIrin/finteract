import openpyxl
import copy
import time

# Configuration
file_name = "FRP_Sched (1).xlsx"
output_file = "FRP_Sched_Merged.xlsx"

print(f"Loading '{file_name}'...")
wb = openpyxl.load_workbook(file_name)

# Create single destination sheet 'Table 1'
out_wb = openpyxl.Workbook()
out_ws = out_wb.active
out_ws.title = "Table 1"

print(f"Total sheets to merge: {len(wb.sheetnames)}")

current_row = 1
total_rows_copied = 0

for i, sheet_name in enumerate(wb.sheetnames, 1):
    src_ws = wb[sheet_name]
    max_r = src_ws.max_row
    max_c = src_ws.max_column
    
    # For Table 1, only take rows 1-52 (rows 53+ were Table 2 manually pasted)
    if sheet_name == "Table 1" and max_r > 52:
        max_r = 52

    # Find last non-empty row in src_ws
    effective_max_r = 0
    for r in range(1, max_r + 1):
        if any(src_ws.cell(r, c).value is not None for c in range(1, max_c + 1)):
            effective_max_r = r

    if effective_max_r == 0:
        continue

    # Copy row by row with complete formatting
    for r in range(1, effective_max_r + 1):
        for c in range(1, max_c + 1):
            src_cell = src_ws.cell(r, c)
            target_cell = out_ws.cell(current_row + r - 1, c)
            target_cell.value = src_cell.value
            
            if src_cell.has_style:
                target_cell.font = copy.copy(src_cell.font)
                target_cell.fill = copy.copy(src_cell.fill)
                target_cell.border = copy.copy(src_cell.border)
                target_cell.alignment = copy.copy(src_cell.alignment)
                target_cell.number_format = src_cell.number_format

    # Copy merged cell ranges with row offset
    for merged_range in src_ws.merged_cells.ranges:
        if merged_range.min_row <= effective_max_r:
            m_max_row = min(merged_range.max_row, effective_max_r)
            new_min_row = current_row + merged_range.min_row - 1
            new_max_row = current_row + m_max_row - 1
            new_min_col = merged_range.min_col
            new_max_col = merged_range.max_col
            
            if new_min_row < new_max_row or new_min_col < new_max_col:
                new_range = f"{openpyxl.utils.get_column_letter(new_min_col)}{new_min_row}:{openpyxl.utils.get_column_letter(new_max_col)}{new_max_row}"
                try:
                    out_ws.merge_cells(new_range)
                except Exception:
                    pass

    # Copy column widths
    for col in src_ws.column_dimensions:
        if src_ws.column_dimensions[col].width:
            curr_w = out_ws.column_dimensions[col].width or 0
            if src_ws.column_dimensions[col].width > curr_w:
                out_ws.column_dimensions[col].width = src_ws.column_dimensions[col].width

    total_rows_copied += effective_max_r
    current_row += effective_max_r

    if i % 20 == 0 or i == len(wb.sheetnames):
        print(f"Processed {i}/{len(wb.sheetnames)} sheets (current row: {current_row - 1})...")

print(f"Saving final merged workbook to '{output_file}'...")
out_wb.save(output_file)
print(f"Merged {len(wb.sheetnames)} sheets successfully into '{output_file}' ({total_rows_copied} total rows)!")
